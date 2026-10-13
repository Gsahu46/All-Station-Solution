"""Writes the roadmap spreadsheet from roadmap.json (formulas recalculated afterwards with the xlsx skill's recalc.py)."""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
W = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(W, "curriculum")); import round5
out = sys.argv[1]
doc = json.load(open(os.path.join(W, "curriculum", "roadmap.json"))); R = doc["stations"]
NEW_ROUND = max(r["round"] for r in R)
F = lambda **k: Font(name="Arial", **k)
HF = PatternFill("solid", fgColor="1D2522"); thin = Border(bottom=Side(style="thin", color="DFE4E1"))
SPEC = PatternFill("solid", fgColor="EEF0FA"); NEWF = PatternFill("solid", fgColor="E6F4F1")
def header(ws, row, cols):
    for c, (h, w) in enumerate(cols, 1):
        cell = ws.cell(row, c, h); cell.font = F(bold=True, color="FFFFFF"); cell.fill = HF
        cell.alignment = Alignment(vertical="center", wrap_text=True); ws.column_dimensions[get_column_letter(c)].width = w
def put(ws, row, vals, fill=None, boldcol=None):
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row, c, v); cell.font = F(size=10, bold=(c == boldcol)); cell.border = thin
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if fill: cell.fill = fill

wb = Workbook(); sm = wb.active; sm.title = "Summary"; st = wb.create_sheet("Stations")
cols = [("ID", 6), ("Stage", 16), ("Suggested experience", 12), ("Track", 14), ("New in v3", 9), ("Title", 36), ("Industry", 15),
        ("Company (fictional)", 14), ("Real-world problem", 55), ("The trap to teach", 45), ("Main endpoints", 40),
        ("Concepts taught", 55), ("Station types", 22), ("Required tests", 60), ("Reviewer view", 9),
        ("Infrastructure", 18), ("Concepts from stations", 14), ("Capstone milestones", 50)]
header(st, 1, cols)
for i, r in enumerate(R, 2):
    new = r["round"] == NEW_ROUND
    vals = [r["id"], r["stage"], r["suggested_experience"], r["track"], "Yes" if new else "No", r["title"], r["industry"], r["company"],
            r["scenario"], r["trap"], "\n".join(r["endpoints"]), "; ".join(r["concepts"]),
            ", ".join(r["station_types"]) or "API only", "\n".join("• " + t for t in r["required_tests"]),
            "Yes" if r["reviewer_view"] else "No", "none" if r["infra"] == "none" else r["infra"].replace("+", ", "),
            ", ".join(map(str, r["prerequisites"])) or "none", "\n".join(r.get("milestones", []))]
    put(st, i, vals, NEWF if new else (SPEC if r["track"] == "Specialisation" else None), boldcol=6)
last = len(R) + 1
st.freeze_panes = "G2"; st.auto_filter.ref = f"A1:R{last}"; st.row_dimensions[1].height = 32

sm.column_dimensions["A"].width = 34
for c in "BCDEF": sm.column_dimensions[c].width = 15
sm["A1"] = doc["name"]; sm["A1"].font = F(bold=True, size=16)
sm["A2"] = doc["note"] + " Version 3: Spring Boot 4, Java 25, after the industry gap research."
sm["A2"].font = F(size=10, italic=True, color="5D6A65")
sm.merge_cells("A2:F2"); sm["A2"].alignment = Alignment(wrap_text=True); sm.row_dimensions[2].height = 48
header(sm, 4, [("Stage", 34), ("Stations", 15), ("New in v3", 15), ("Optional", 15), ("Reviewer view", 15), ("Share", 15)])
rs = f"Stations!$B$2:$B${last}"; rn = f"Stations!$E$2:$E${last}"; rt = f"Stations!$D$2:$D${last}"; rv = f"Stations!$O$2:$O${last}"
for i, (stg, yrs) in enumerate([("Foundations", "0-1"), ("Working Systems", "0-2"), ("Production Scale", "2-3")], 5):
    sm.cell(i, 1, f"{stg} (suggested {yrs} years)").font = F()
    for c, f in ((2, f'=COUNTIF({rs},"{stg}")'), (3, f'=COUNTIFS({rs},"{stg}",{rn},"Yes")'),
                 (4, f'=COUNTIFS({rs},"{stg}",{rt},"Specialisation")'), (5, f'=COUNTIFS({rs},"{stg}",{rv},"Yes")')):
        sm.cell(i, c, f).font = F()
    x = sm.cell(i, 6, f"=IF($B$8=0,0,B{i}/$B$8)"); x.font = F(); x.number_format = "0.0%"
sm.cell(8, 1, "Total").font = F(bold=True)
for c in "BCDE": sm[f"{c}8"] = f"=SUM({c}5:{c}7)"; sm[f"{c}8"].font = F(bold=True)
x = sm.cell(8, 6, "=IF($B$8=0,0,B8/$B$8)"); x.font = F(bold=True); x.number_format = "0.0%"
header(sm, 10, [("Station type", 34), ("Stations", 15)])
for i, t in enumerate(["Database", "Security", "Concurrency", "Messaging", "Idempotency", "Operations"], 11):
    sm.cell(i, 1, t).font = F(); sm.cell(i, 2, f'=COUNTIF(Stations!$M$2:$M${last},"*{t}*")').font = F()
sm["A18"] = "New stations are shaded green and optional ones blue on the Stations sheet. A station can have several types."
sm["A18"].font = F(size=9, italic=True, color="5D6A65")

ch = wb.create_sheet("Changes in v3")
header(ch, 1, [("Change", 12), ("Station", 46), ("Detail", 80)])
row = 2
title_id = {r["title"]: r["id"] for r in R}
for n in round5.NEW5:
    put(ch, row, ["Added", f'{title_id[n[2]]:03d} {n[2]}', "; ".join(n[7])], NEWF, boldcol=2); row += 1
put(ch, row, ["Replaced", f'{title_id["Passkey sign-in with WebAuthn"]:03d} Passkey sign-in with WebAuthn', "Replaces 'Two-factor login with TOTP'; TOTP stays as the fallback"], NEWF, boldcol=2); row += 1
for old, new in round5.RENAMES.items():
    if old != "Two-factor login with TOTP":
        put(ch, row, ["Renamed", f'{title_id[new]:03d} {new}', f"Was '{old}'"], boldcol=2); row += 1
for gone, (target, extra) in round5.MERGES.items():
    t = round5.RENAMES.get(target, target)
    put(ch, row, ["Merged", gone, f"Now part of {title_id[t]:03d} {t}" + (f" ({'; '.join(extra)})" if extra else "")], boldcol=2); row += 1
for t, u in round5.UPDATES.items():
    if t in round5.RENAMES.values():
        continue
    what = "; ".join(u.get("concepts+", [])) or ("concepts: " + "; ".join(u["concepts"]) if "concepts" in u else "") or ("endpoints: " + "; ".join(u.get("endpoints", [])))
    put(ch, row, ["Updated", f"{title_id[t]:03d} {t}", what], boldcol=2); row += 1
ch.freeze_panes = "A2"; ch.auto_filter.ref = f"A1:C{row - 1}"

ss = wb.create_sheet("Station standard")
header(ss, 1, [("README section and rules", 120)])
for i, line in enumerate([l for l in open(os.path.join(W, "STATION_STANDARD.md")).read().splitlines() if l.strip()], 2):
    put(ss, i, [line.replace("**", "").lstrip("# ")])

cm = wb.create_sheet("Capstone milestones")
header(cm, 1, [("Capstone", 38), ("Milestone", 80)])
row = 2
for r in R:
    for m in r.get("milestones", []):
        put(cm, row, [f'{r["id"]:03d} {r["title"]}', m], boldcol=1); row += 1
wb.save(out); print("saved", out)
